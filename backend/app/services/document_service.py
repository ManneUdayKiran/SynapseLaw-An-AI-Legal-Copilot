import re
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from fastapi import HTTPException, UploadFile, status
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.db.models import Document, User
from app.rag.pipeline import index_document
from app.utils.file_validation import validate_upload


from app.core.security import strict_sanitize_contract_text


def clean_text(text: str) -> str:
    return strict_sanitize_contract_text(text)


def extract_text(path: Path, extension: str) -> str:
    if extension == ".txt":
        return clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    if extension == ".pdf":
        parts: list[str] = []
        try:
            with fitz.open(path) as pdf:
                for index, page in enumerate(pdf, start=1):
                    page_text = clean_text(page.get_text("text"))
                    if page_text:
                        parts.append(f"[Page {index}]\n{page_text}")
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read PDF text") from exc
        return clean_text("\n\n".join(parts))
    if extension == ".docx":
        try:
            doc = DocxDocument(path)
            return clean_text("\n".join(p.text for p in doc.paragraphs if p.text.strip()))
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read DOCX text") from exc
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported file type")


async def upload_document(db: Session, owner: User, file: UploadFile) -> Document:
    settings = get_settings()
    validated, data = await validate_upload(file, settings.max_upload_bytes)
    existing = db.scalar(
        select(Document).where(Document.owner_id == owner.id, Document.sha256 == validated.sha256)
    )
    if existing:
        return existing

    storage_root = settings.storage_path / owner.id
    storage_root.mkdir(parents=True, exist_ok=True)
    stored_name = f"{validated.sha256}{validated.extension}"
    stored_path = (storage_root / stored_name).resolve()
    if storage_root not in stored_path.parents:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid storage path")
    stored_path.write_bytes(data)

    extracted = extract_text(stored_path, validated.extension)
    if not extracted:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Document contains no readable text")

    document = Document(
        owner_id=owner.id,
        filename=Path(file.filename or "document").name,
        stored_name=stored_name,
        content_type=file.content_type or "application/octet-stream",
        size_bytes=validated.size_bytes,
        sha256=validated.sha256,
        extracted_text=extracted,
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    chunks = index_document(document.id, extracted)
    document.chunk_count = len(chunks)
    db.commit()
    db.refresh(document)
    return document


def get_owned_document(db: Session, owner: User, document_id: str) -> Document:
    document = db.get(Document, document_id)
    if document is None or document.owner_id != owner.id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    return document


def list_documents(db: Session, owner: User) -> list[Document]:
    return list(
        db.scalars(select(Document).where(Document.owner_id == owner.id).order_by(Document.created_at.desc()))
    )
